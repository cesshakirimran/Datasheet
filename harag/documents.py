"""Document loading and chunking.

Supported formats: .md, .txt, .pdf (pypdf), .docx (python-docx).
Markdown headings are tracked so every chunk knows its section, which is shown in citations.
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

SUPPORTED = {".md", ".txt", ".pdf", ".docx"}


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str          # file name, e.g. hr_leave_policy.md
    title: str           # document title (first heading or file name)
    section: str         # nearest heading above the chunk
    text: str

    def to_dict(self) -> dict:
        return asdict(self)

    @staticmethod
    def from_dict(d: dict) -> "Chunk":
        return Chunk(**d)


# ---------------------------------------------------------------- loading

def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        import docx  # python-docx

        d = docx.Document(str(path))
        lines = []
        for p in d.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            text = p.text.strip()
            if not text:
                continue
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "2"
                lines.append("#" * int(level) + " " + text)
            else:
                lines.append(text)
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(lines)
    raise ValueError(f"Unsupported file type: {path}")


def iter_documents(corpus_dir: Path) -> Iterable[tuple[str, str]]:
    for path in sorted(corpus_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            yield str(path.relative_to(corpus_dir)), read_document(path)


# ---------------------------------------------------------------- chunking

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def _blocks_with_sections(text: str) -> tuple[str, list[tuple[str, str]]]:
    """Split a document into (section, block) pairs; blocks are paragraphs, list items or table rows."""
    title = ""
    section = ""
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append((section, " ".join(paragraph).strip()))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        m = _HEADING.match(line)
        if m:
            flush()
            heading = m.group(2).strip()
            if not title:
                title = heading
            section = heading
            continue
        if not line:
            flush()
            continue
        if re.match(r"^\|?\s*-{3,}", line):  # markdown table separator
            continue
        if line.startswith(("- ", "* ", "|")) or re.match(r"^\d+\.\s", line):
            flush()
            blocks.append((section, line.strip("| ").replace(" | ", "; ")))
            continue
        paragraph.append(line)
    flush()
    return title, blocks


def chunk_document(doc_id: str, text: str, chunk_size: int = 180, overlap: int = 30) -> list[Chunk]:
    """Greedy chunking that keeps sections together and never splits inside a block if avoidable.

    Chunks are prefixed with the document title and section so that the retriever and the
    generator see where the text comes from.
    """
    title, blocks = _blocks_with_sections(text)
    title = title or Path(doc_id).stem
    chunks: list[Chunk] = []
    current: list[str] = []
    current_section = blocks[0][0] if blocks else ""

    def emit() -> None:
        if not current:
            return
        body = "\n".join(current)
        header = f"{title} > {current_section}" if current_section and current_section != title else title
        cid = hashlib.sha1(f"{doc_id}|{len(chunks)}|{body}".encode()).hexdigest()[:12]
        chunks.append(Chunk(cid, doc_id, title, current_section, f"[{header}]\n{body}"))

    for section, block in blocks:
        words_now = sum(len(b.split()) for b in current)
        if current and (section != current_section or words_now + len(block.split()) > chunk_size):
            emit()
            if section == current_section and overlap > 0:
                # carry the tail of the previous chunk as overlap
                tail_words = " ".join(current).split()[-overlap:]
                current = [" ".join(tail_words)]
            else:
                current = []
            current_section = section
        current.append(block)
    emit()
    return chunks


def build_chunks(corpus_dir: Path, chunk_size: int, overlap: int) -> list[Chunk]:
    chunks: list[Chunk] = []
    for doc_id, text in iter_documents(corpus_dir):
        chunks.extend(chunk_document(doc_id, text, chunk_size, overlap))
    if not chunks:
        raise FileNotFoundError(f"No documents ({', '.join(sorted(SUPPORTED))}) found in {corpus_dir}")
    return chunks


def save_chunks(chunks: list[Chunk], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c.to_dict(), ensure_ascii=False) + "\n")


def load_chunks(path: Path) -> list[Chunk]:
    with path.open(encoding="utf-8") as f:
        return [Chunk.from_dict(json.loads(line)) for line in f if line.strip()]
